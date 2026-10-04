import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "output"
stores = pd.read_csv(f"{OUT}/stores.csv")
products = pd.read_csv(f"{OUT}/products.csv")
sales = pd.read_csv(f"{OUT}/weekly_sales.csv", keep_default_na=False)

F = "Arial"
H_FILL = PatternFill("solid", start_color="1F3A5F")
H_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
INPUT = Font(name=F, size=10, color="0000FF")
KEY = PatternFill("solid", start_color="FFFF00")
TITLE = Font(name=F, bold=True, size=14, color="1F3A5F")
thin = Side(style="thin", color="D0D0D0")
RS = '"₹"#,##0;("₹"#,##0);-'
RS2 = '"₹"#,##0.00'

wb = Workbook()

# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
lines = [
    ("WrapStreet Pricing Prototype - Synthetic Dataset", TITLE),
    ("Fictional 8-outlet Mumbai QSR chain | 6 products | 104 weeks (Apr 2024 - Mar 2026) | 4,992 rows", BODY),
    ("", BODY),
    ("Sheets", Font(name=F, bold=True, size=11)),
    ("Assumptions - every global parameter used to generate and evaluate the data (blue = editable input)", BODY),
    ("Stores - outlet profiles: type, income tier, competition, footfall, delivery share", BODY),
    ("Products - menu, current prices (ex-GST), unit food + packaging cost", BODY),
    ("Weekly_Sales - one row per store x product x week; grey-header columns are live formulas", BODY),
    ("", BODY),
    ("Important", Font(name=F, bold=True, size=11)),
    ("All data is SYNTHETIC. Store areas are real Mumbai neighbourhoods; the chain, prices and volumes are invented.", BODY),
    ("The planted 'true' elasticities are kept OUT of this workbook (planted_truth.csv) so the model must estimate them.", BODY),
    ("Last 8 weeks (week 97-104, Feb-Mar 2026) = 'current state' baseline: no tests or discounts.", BODY),
    ("Contribution = net revenue - food/packaging cost - aggregator commission on delivery orders. Fixed costs excluded.", BODY),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = f
ws.column_dimensions["A"].width = 110

# ---------------------------------------------------------------- Assumptions
wa = wb.create_sheet("Assumptions")
assump = [
    ("Parameter", "Value", "Unit", "Rationale / source"),
    ("GST rate (excluded from revenue)", 0.05, "%", "Restaurant GST is a pass-through; revenue is shown ex-GST"),
    ("Aggregator commission", 0.25, "% of order value", "Assumed typical commission on Swiggy/Zomato delivery orders"),
    ("Monsoon delivery-share shift", 0.10, "pp", "Assumption: more delivery during Jun-Sep rains"),
    ("Noise (std dev of log orders)", 0.08, "log units", "Assumption: ~8% random week-to-week variation"),
    ("Organic growth", 0.03, "% per year", "Assumption"),
    ("Promo visibility uplift", 0.05, "%", "Assumption: discount weeks get extra visibility on top of the price effect"),
    ("FY24-25 price level vs today", 0.06, "% lower", "Assumption: one chain-wide revision in Apr 2025"),
    ("Competitor price drift", 0.05, "% per year", "Assumption"),
    ("Competitor sensitivity - High competition", 0.5, "elasticity", "Assumption: share of demand moving with competitor price"),
    ("Competitor sensitivity - Medium competition", 0.3, "elasticity", "Assumption"),
    ("Baseline window", 8, "weeks", "Last 8 weeks used as 'current state'"),
    ("Random seed", 42, "-", "Re-running generate_data.py reproduces this exact dataset"),
]
for r, row in enumerate(assump, 1):
    for c, v in enumerate(row, 1):
        cell = wa.cell(row=r, column=c, value=v)
        if r == 1:
            cell.font, cell.fill = H_FONT, H_FILL
        else:
            cell.font = INPUT if c == 2 else BODY
            if c == 2 and isinstance(v, float) and v < 1:
                cell.number_format = "0.0%" if row[2].startswith("%") or row[2] == "pp" else "0.00"
wa["B3"].fill = KEY  # commission drives contribution formulas
for col, w in zip("ABCD", (42, 10, 18, 80)):
    wa.column_dimensions[col].width = w
wa.freeze_panes = "A2"
COMM = "Assumptions!$B$3"


def write_table(ws, df, fmts=None, formula_cols=None):
    fmts = fmts or {}
    formula_cols = formula_cols or {}
    cols = list(df.columns) + list(formula_cols.keys())
    for c, name in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = H_FONT
        cell.fill = PatternFill("solid", start_color="5A6B7D") if name in formula_cols else H_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r, rec in enumerate(df.itertuples(index=False), 2):
        for c, v in enumerate(rec, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY
            if df.columns[c - 1] in fmts:
                cell.number_format = fmts[df.columns[c - 1]]
        for k, (name, tmpl) in enumerate(formula_cols.items(), len(df.columns) + 1):
            cell = ws.cell(row=r, column=k, value=tmpl.format(r=r))
            cell.font = BODY
            if name in fmts:
                cell.number_format = fmts[name]
    for c, name in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(c)].width = max(11, min(28, len(name) + 3))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# ---------------------------------------------------------------- Stores / Products
write_table(wb.create_sheet("Stores"), stores,
            {"delivery_share": "0%", "competitor_price_ratio": "0.00", "footfall_index": "0.00"})
write_table(wb.create_sheet("Products"), products,
            {"current_price": RS, "unit_cost": RS})

# ---------------------------------------------------------------- Weekly_Sales
raw = sales[["week_start", "fy", "week_no", "store_id", "product_id", "menu_price",
             "discount_pct", "competitor_price", "price_test_flag", "promo_flag",
             "festival_flag", "festival_name", "season_index", "orders", "delivery_orders"]].copy()
raw["week_start"] = pd.to_datetime(raw.week_start)
# raw cols A..O ; formulas P..U
formulas = {
    "effective_price": "=F{r}*(1-G{r})",
    "dine_in_orders": "=N{r}-O{r}",
    "unit_cost": "=INDEX(Products!$E:$E,MATCH(E{r},Products!$A:$A,0))",
    "net_revenue": "=N{r}*P{r}",
    "variable_cost": "=N{r}*R{r}+O{r}*P{r}*" + COMM,
    "contribution": "=S{r}-T{r}",
}
write_table(wb.create_sheet("Weekly_Sales"), raw,
            {"week_start": "dd-mmm-yy", "menu_price": RS, "discount_pct": "0%;;-",
             "competitor_price": RS, "effective_price": RS2, "unit_cost": RS,
             "net_revenue": RS, "variable_cost": RS, "contribution": RS},
            formulas)

wb.save(f"{OUT}/wrapstreet_synthetic_data.xlsx")
print("saved")
